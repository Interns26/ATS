from io import BytesIO
from pathlib import Path
from pypdf import PdfReader
from docx import Document


def extract_text_from_bytes(file_bytes: bytes, ext: str) -> str:
    """Extract text from in-memory file bytes (no disk write needed).

    Supports .pdf, .docx, and .txt — used for job description uploads,
    which arrive as bytes over HTTP rather than as a path on disk.
    """
    ext = ext.lower()

    if ext == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        text = ""
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"
        return text

    elif ext == ".docx":
        doc = Document(BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)

    elif ext == ".txt":
        return file_bytes.decode("utf-8", errors="ignore")

    else:
        raise ValueError("Unsupported file format.")


def extract_pdf(path: str):

    reader = PdfReader(path)

    text = ""

    for page in reader.pages:
        text += page.extract_text() + "\n"

    return text


def extract_docx(path: str):

    doc = Document(path)

    return "\n".join(p.text for p in doc.paragraphs)


def extract_resume_text(path: str):

    ext = Path(path).suffix.lower()

    if ext == ".pdf":
        return extract_pdf(path)

    elif ext == ".docx":
        return extract_docx(path)

    else:
        raise ValueError("Unsupported file format.")